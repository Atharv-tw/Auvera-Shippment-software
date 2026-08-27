"""Export reconciled tracker rows to an .xlsx in the Shipment-01-04-2026 layout."""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string

from app.services.tracker_map import TRACKER_COLUMNS, TYPE_BY_KEY

SHEET_TITLE = "Shipment-01-04-2026"
_HEADER_ROW = 2
_DATA_START = 3
_DATE_FMT = "dd-mmm-yy"

# columns that get a Sub Total in row 1 (mirrors the source workbook)
_SUBTOTAL_KEYS = {"order_qty", "ship_qty", "buyer_total_value", "vendor_total_value"}


def _as_date(value):
    if isinstance(value, (date, datetime)):
        return value
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value).date()
        except ValueError:
            return value
    return value


def build_tracker_workbook(rows, keys: list[str] | None = None) -> Workbook:
    """The tracker as a workbook.

    With no ``keys`` this is the full 56-column sheet, each column at its
    canonical letter, which is what the shipping team's own file looks like.
    Given ``keys`` it exports just those columns packed contiguously from A -
    a filtered export should read like a normal sheet, not the full template
    with holes punched in it.
    """
    subset = keys is not None
    columns = (
        [c for c in TRACKER_COLUMNS if c["key"] in set(keys)] if subset else TRACKER_COLUMNS
    )
    # position -> Excel column index; contiguous when a subset was asked for
    index_of = {
        c["key"]: (i + 1 if subset else column_index_from_string(c["col"]))
        for i, c in enumerate(columns)
    }

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE

    header_font = Font(bold=True, size=9)
    header_fill = PatternFill("solid", fgColor="E2E8F0")
    wrap = Alignment(wrap_text=True, vertical="center")

    # header row
    for c in columns:
        cell = ws.cell(row=_HEADER_ROW, column=index_of[c["key"]], value=c["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    subtotals: dict[str, float] = {k: 0.0 for k in _SUBTOTAL_KEYS if k in index_of}

    # data rows
    for i, row in enumerate(rows):
        r = _DATA_START + i
        data = row.data if hasattr(row, "data") else row.get("data", {})
        data = data or {}
        for col in columns:
            key = col["key"]
            val = data.get(key)
            if val is None or val == "":
                continue
            idx = index_of[key]
            if TYPE_BY_KEY[key] == "date":
                dv = _as_date(val)
                cell = ws.cell(row=r, column=idx, value=dv)
                if isinstance(dv, (date, datetime)):
                    cell.number_format = _DATE_FMT
            elif TYPE_BY_KEY[key] == "number":
                try:
                    num = float(val)
                except (TypeError, ValueError):
                    ws.cell(row=r, column=idx, value=val)
                    continue
                ws.cell(row=r, column=idx, value=num)
                if key in subtotals:
                    subtotals[key] += num
            else:
                ws.cell(row=r, column=idx, value=str(val))

    # sub-total row
    ws.cell(row=1, column=1, value="Sub Total").font = Font(bold=True, size=9)
    for key, total in subtotals.items():
        ws.cell(row=1, column=index_of[key], value=round(total, 4)).font = Font(bold=True, size=9)

    ws.freeze_panes = "A3"
    return wb


def tracker_workbook_bytes(rows, keys: list[str] | None = None) -> bytes:
    wb = build_tracker_workbook(rows, keys)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
