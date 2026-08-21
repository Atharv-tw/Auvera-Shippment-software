"""Reporting endpoints for the CEO and admin.

One summary shape, sliced by season / customer / vendor, plus the same figures
as a downloadable workbook.
"""

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Response, status
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.database import get_db
from app.dependencies import require_reports
from app.models import User
from app.schemas import ReportBucket, ReportOut
from app.services import reports as rp

router = APIRouter(prefix="/api/reports", tags=["reports"])

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# (schema field, sheet heading)
_COLUMNS = [
    ("label", "Group"),
    ("po_count", "POs"),
    ("line_count", "Lines"),
    ("order_qty", "Order qty"),
    ("ship_qty", "Ship qty"),
    ("short_extra_qty", "Short / Extra"),
    ("buyer_value", "Buyer value"),
    ("vendor_value", "Vendor value"),
    ("margin", "Margin"),
    ("margin_pct", "Margin %"),
    ("shipped_lines", "Shipped lines"),
    ("pending_lines", "Pending lines"),
    ("avg_shipment_delay", "Avg shipment delay"),
    ("avg_docs_delay", "Avg docs delay"),
]

_GROUP_TITLES = {"season": "Season", "customer": "Customer", "vendor": "Vendor"}


def _build(db: Session, group_by: str, season, customer, vendor, date_from, date_to) -> dict:
    try:
        return rp.build_report(
            db, group_by=group_by, season=season, customer=customer,
            vendor=vendor, date_from=date_from, date_to=date_to,
        )
    except ValueError as exc:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, str(exc))


@router.get("/options")
def options(db: Session = Depends(get_db), _user: User = Depends(require_reports)):
    """Seasons, customers and vendors that actually have data behind them."""
    return rp.filter_options(db)


@router.get("/summary", response_model=ReportOut)
def summary(
    group_by: str = "season",
    season: str | None = None,
    customer: str | None = None,
    vendor: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    _user: User = Depends(require_reports),
):
    report = _build(db, group_by, season, customer, vendor, date_from, date_to)
    return ReportOut(
        group_by=report["group_by"],
        buckets=[ReportBucket(**b) for b in report["buckets"]],
        totals=ReportBucket(**report["totals"]),
        filters=report["filters"],
    )


@router.get("/export")
def export_report(
    group_by: str = "season",
    season: str | None = None,
    customer: str | None = None,
    vendor: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
    db: Session = Depends(get_db),
    _user: User = Depends(require_reports),
):
    report = _build(db, group_by, season, customer, vendor, date_from, date_to)
    title = _GROUP_TITLES.get(group_by, "Report")

    wb = Workbook()
    ws = wb.active
    ws.title = f"By {title}"

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E2E8F0")

    ws.cell(row=1, column=1, value=f"Report by {title.lower()}").font = Font(bold=True, size=12)
    active = [f"{k}: {v}" for k, v in report["filters"].items() if v]
    ws.cell(row=2, column=1, value="Filters - " + ("; ".join(active) if active else "none"))

    header_row = 4
    for index, (_key, heading) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=index, value=heading)
        cell.font = bold
        cell.fill = fill
    ws.cell(row=header_row, column=1, value=title).font = bold

    for offset, bucket in enumerate(report["buckets"], start=header_row + 1):
        for index, (key, _heading) in enumerate(_COLUMNS, start=1):
            ws.cell(row=offset, column=index, value=bucket.get(key))

    total_row = header_row + len(report["buckets"]) + 1
    ws.cell(row=total_row, column=1, value="TOTAL").font = bold
    for index, (key, _heading) in enumerate(_COLUMNS, start=1):
        if key == "label":
            continue
        ws.cell(row=total_row, column=index, value=report["totals"].get(key)).font = bold

    ws.freeze_panes = f"A{header_row + 1}"
    for index in range(1, len(_COLUMNS) + 1):
        ws.column_dimensions[ws.cell(row=header_row, column=index).column_letter].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"{title} report - {datetime.now():%d-%m-%Y}.xlsx"
    return Response(
        content=buf.getvalue(),
        media_type=_XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
