"""Exporting the current view rather than the whole table.

If a filtered, narrowed view cannot be exported as it appears, people rebuild
the sheet by hand in Excel - which is the behaviour this whole change set is
trying to remove.
"""

import io

from openpyxl import load_workbook

from .conftest import CUSTOMER_FILES, SAMPLES

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _auth(token):
    return {"Authorization": f"Bearer {token}"}


def _admin_with_rows(client):
    r = client.post("/api/auth/register", json={
        "email": "admin@example.com", "password": "secret123", "name": "admin", "role": "admin",
    })
    token = r.json()["access_token"]
    path = SAMPLES / CUSTOMER_FILES[0]
    files = [("files", (path.name, path.read_bytes(), XLSX))]
    assert client.post("/api/orders/upload", files=files, headers=_auth(token)).status_code == 200
    return token


def _sheet(content: bytes):
    return load_workbook(io.BytesIO(content)).active


def test_export_without_a_view_is_still_the_full_sheet(client):
    token = _admin_with_rows(client)
    r = client.post("/api/tracker/export", json={}, headers=_auth(token))
    assert r.status_code == 200
    ws = _sheet(r.content)
    # header row carries all 56 columns at their canonical letters
    assert ws["E2"].value == "Buyer PO#"
    assert ws["G2"].value == "Style No."


def test_export_narrows_to_the_visible_columns_packed_from_A(client):
    token = _admin_with_rows(client)
    r = client.post(
        "/api/tracker/export",
        json={"keys": ["buyer_po", "style_no", "vendor_total_value"]},
        headers=_auth(token),
    )
    assert r.status_code == 200
    ws = _sheet(r.content)
    headers = [ws.cell(row=2, column=i).value for i in range(1, 5)]
    # contiguous from A - a filtered export should read like a normal sheet,
    # not the full template with holes in it
    assert headers[:3] == ["Buyer PO#", "Style No.", "Vendor Total Value"]
    assert headers[3] is None


def test_export_keeps_only_the_requested_rows_in_the_given_order(client):
    token = _admin_with_rows(client)
    rows = client.get("/api/tracker", headers=_auth(token)).json()
    assert len(rows) >= 2
    wanted = [rows[1]["id"], rows[0]["id"]]  # deliberately reversed

    r = client.post(
        "/api/tracker/export",
        json={"row_ids": wanted, "keys": ["buyer_po", "style_no"]},
        headers=_auth(token),
    )
    ws = _sheet(r.content)
    exported = [
        (ws.cell(row=i, column=1).value, ws.cell(row=i, column=2).value)
        for i in range(3, 3 + len(wanted))
    ]
    expected = [
        (rows[1]["buyer_po"], rows[1]["style_no"]),
        (rows[0]["buyer_po"], rows[0]["style_no"]),
    ]
    assert exported == expected
    assert ws.cell(row=3 + len(wanted), column=1).value is None, "no extra rows"
